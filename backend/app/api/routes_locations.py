"""
Location management endpoints.

POST   /api/locations              – save a farm/AOI coordinate
GET    /api/locations              – list all saved locations
GET    /api/locations/{id}         – get a single location
PATCH  /api/locations/{id}         – update status / link job IDs
DELETE /api/locations/{id}         – delete a location
GET    /api/locations/{id}/excel   – download Excel summary

All write operations are persisted to the database configured via
DATABASE_URL in backend/.env (default: SQLite at data/agrisense.db).
"""
from __future__ import annotations

import io
from typing import Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field

from app.services.location_service import location_service

router = APIRouter()


# ---------------------------------------------------------------------------
# Request / Response schemas
# ---------------------------------------------------------------------------

class CreateLocationRequest(BaseModel):
    latitude:         float = Field(..., ge=-90,   le=90)
    longitude:        float = Field(..., ge=-180,  le=180)
    name:             str   = Field("", max_length=120)
    radius_km:        float = Field(10.0, ge=0.1,  le=500)
    date_from:        str   = ""
    date_to:          str   = ""
    cloud_threshold:  float = Field(20.0, ge=0,    le=100)
    ndvi_threshold:   float = Field(0.25, ge=0,    le=1)
    user_id:          str   = "default"


class UpdateLocationRequest(BaseModel):
    status:          Optional[str]  = None
    analyze_job_id:  Optional[str]  = None
    classify_job_id: Optional[str]  = None
    result_summary:  Optional[dict] = None


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

@router.post("/locations", status_code=201)
async def create_location(request: CreateLocationRequest):
    """Save a new farm / AOI coordinate to the database."""
    loc = await location_service.create(
        latitude=request.latitude,
        longitude=request.longitude,
        name=request.name,
        radius_km=request.radius_km,
        date_from=request.date_from,
        date_to=request.date_to,
        cloud_threshold=request.cloud_threshold,
        ndvi_threshold=request.ndvi_threshold,
        user_id=request.user_id,
    )
    return loc.to_dict()


@router.get("/locations")
async def list_locations(user_id: Optional[str] = None):
    """List all saved locations from the database (newest first)."""
    locs = await location_service.list_all(user_id=user_id)
    return {
        "locations": [l.to_dict() for l in locs],
        "count":     len(locs),
    }


@router.get("/locations/{location_id}")
async def get_location(location_id: str):
    """Get a single saved location by ID."""
    loc = await location_service.get(location_id)
    if loc is None:
        raise HTTPException(status_code=404, detail=f"Location '{location_id}' not found.")
    return loc.to_dict()


@router.patch("/locations/{location_id}")
async def update_location(location_id: str, request: UpdateLocationRequest):
    """Update location status or link analysis / classify job IDs."""
    loc = await location_service.update_status(
        location_id=location_id,
        status=request.status or "",
        analyze_job_id=request.analyze_job_id,
        classify_job_id=request.classify_job_id,
        result_summary=request.result_summary,
    )
    if loc is None:
        raise HTTPException(status_code=404, detail=f"Location '{location_id}' not found.")
    return loc.to_dict()


@router.delete("/locations/{location_id}", status_code=204)
async def delete_location(location_id: str):
    """Delete a saved location from the database."""
    deleted = await location_service.delete(location_id)
    if not deleted:
        raise HTTPException(status_code=404, detail=f"Location '{location_id}' not found.")
    return None


@router.get("/locations/{location_id}/excel")
async def export_location_excel(location_id: str):
    """
    Download an Excel file summarising the saved location and any
    linked analysis results.
    """
    loc = await location_service.get(location_id)
    if loc is None:
        raise HTTPException(status_code=404, detail=f"Location '{location_id}' not found.")

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Location Summary"

        hdr_font = Font(bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")

        for col, h in enumerate(["Field", "Value"], 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center")

        d = loc.to_dict()
        rows = [
            ("Name",              d["name"]),
            ("Latitude",          d["latitude"]),
            ("Longitude",         d["longitude"]),
            ("Radius (km)",       d["radius_km"]),
            ("Date From",         d["date_from"]),
            ("Date To",           d["date_to"]),
            ("Cloud Threshold %", d["cloud_threshold"]),
            ("NDVI Threshold",    d["ndvi_threshold"]),
            ("User ID",           d["user_id"]),
            ("Status",            d["status"]),
            ("Analyze Job ID",    d["analyze_job_id"] or ""),
            ("Classify Job ID",   d["classify_job_id"] or ""),
            ("Saved At",          d["created_at"] or ""),
            ("Last Updated",      d["updated_at"] or ""),
        ]

        if d.get("result_summary"):
            rs = d["result_summary"]
            rows += [
                ("─── Analysis Results ───", ""),
                ("Total Area (ha)",           rs.get("total_area_ha", "")),
                ("Vegetated Area (ha)",        rs.get("vegetated_area_ha", "")),
                ("Predicted Brinjal (ha)",     rs.get("predicted_brinjal_ha", "")),
                ("Brinjal % of Area",          rs.get("predicted_brinjal_pct", "")),
                ("Avg Brinjal Confidence",     rs.get("avg_brinjal_confidence", "")),
                ("Num Brinjal Fields",         rs.get("num_brinjal_fields", "")),
            ]

        for r_idx, (field, value) in enumerate(rows, 2):
            ws.cell(row=r_idx, column=1, value=field)
            ws.cell(row=r_idx, column=2, value=value)

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 35

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        return Response(
            content=buf.read(),
            media_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f"attachment; filename=location_{location_id[:8]}.xlsx"
                )
            },
        )
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="openpyxl is not installed. Add it to requirements.txt.",
        )
